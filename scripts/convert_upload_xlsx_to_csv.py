from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

SOURCE = PROJECT_ROOT / "data" / "사용자별_더미문서_업로드목록.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "csv"


def safe_sheet_name(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*]', "_", name).strip() or "sheet"


def normalize_cell(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    report = []

    for sheet in [workbook["업로드목록"]]:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [normalize_cell(value) for value in row]
            if any(str(value) != "" for value in values):
                rows.append(values)

        output_path = OUTPUT_DIR / f"{safe_sheet_name(sheet.title)}.csv"
        with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerows(rows)

        with output_path.open("r", newline="", encoding="utf-8-sig") as handle:
            roundtrip_rows = list(csv.reader(handle))

        original_as_text = [[str(value) for value in row] for row in rows]
        report.append(
            {
                "sheet": sheet.title,
                "csv": str(output_path),
                "rows": len(rows),
                "cols": max((len(row) for row in rows), default=0),
                "roundtrip_equal_as_text": original_as_text == roundtrip_rows,
            }
        )

    workbook.close()
    report_path = OUTPUT_DIR / "conversion_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
